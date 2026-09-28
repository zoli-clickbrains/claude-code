"""Parancssor.

    python -m icp tablazat ICP.xlsx [--minta]   # táblázat készítése (mintaadatokkal)
    python -m icp pontoz ugyfelek.xlsx|.csv [--csv eredmeny.csv]
"""

import argparse
import csv
import sys

from openpyxl import load_workbook

from .config import INPUT_COLUMNS
from .engine import score_clients
from .sample import SAMPLE_CLIENTS
from .workbook import S_IN, build_workbook

HEADER_TO_KEY = {header: key for key, header, *_ in INPUT_COLUMNS}


def read_clients(path):
    """Ügyfelek beolvasása az Ügyfelek lapról (xlsx) vagy CSV-ből; a fejlécek a táblázatéval egyeznek."""
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as fh:
            rows = list(csv.reader(fh))
    else:
        ws = load_workbook(path, data_only=True)[S_IN]
        rows = [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
    if not rows:
        return []
    keys = [HEADER_TO_KEY.get(str(h).strip()) for h in rows[0]]
    unknown = [h for h, k in zip(rows[0], keys) if k is None and str(h).strip()]
    if unknown:
        raise SystemExit(f"Ismeretlen oszlop(ok): {', '.join(map(str, unknown))}")
    return [{k: v for k, v in zip(keys, row) if k} for row in rows[1:]]


def _num(value):
    if isinstance(value, str) and value.strip():
        return float(value.replace(" ", "").replace(",", "."))
    return value


def main(argv=None):
    parser = argparse.ArgumentParser(prog="icp", description="ICP ügyfélminősítő")
    sub = parser.add_subparsers(dest="cmd", required=True)
    t = sub.add_parser("tablazat", help="Excel/Google Táblázat sablon készítése")
    t.add_argument("kimenet")
    t.add_argument("--minta", action="store_true", help="kitalált mintaügyfelekkel töltse fel")
    s = sub.add_parser("pontoz", help="ügyfelek pontozása és rangsorolása")
    s.add_argument("bemenet")
    s.add_argument("--csv", help="eredmény mentése CSV-be")
    args = parser.parse_args(argv)

    if args.cmd == "tablazat":
        build_workbook(SAMPLE_CLIENTS if args.minta else ()).save(args.kimenet)
        print(f"Elkészült: {args.kimenet}")
        return 0

    text_keys = {key for key, _, _, kind, _ in INPUT_COLUMNS if kind in ("szoveg", "lista", "igennem")}
    clients = [{k: (v if k in text_keys else _num(v)) for k, v in c.items()} for c in read_clients(args.bemenet)]
    results = score_clients(clients)
    fields = ["rang", "nev", "pontszam", "rendszer", "statusz", "statusz_nev", "dontes", "kizart", "javasolt_dij",
              "jelenlegi_dij", "elteres", "dij_alapja", "jutalekos_modell", "jutalekos_fix_dij", "jutalek_szazalek",
              "legerosebb", "leggyengebb", "hianyzo_adatok"]
    if args.csv:
        with open(args.csv, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            groups = list(results[0].csoportok) if results else []
            w.writerow(fields + groups)
            for r in results:
                w.writerow([getattr(r, f) for f in fields] + [r.csoportok[g] for g in groups])
        print(f"Elmentve: {args.csv}")
    for r in results:
        fee = "—" if r.javasolt_dij is None else f"{r.javasolt_dij:,} Ft".replace(",", " ")
        print(f"{r.rang:>3}. {r.nev:<34} ICP {r.pontszam:5.1f}  rendszer {r.rendszer:5.1f}  "
              f"{r.statusz}. {r.statusz_nev:<42} {r.dontes:<32} díj: {fee}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
