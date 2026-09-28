"""Az ICP-ügyfélminősítő táblázat előállítása élő képletekkel.

A képletek Excelben, Google Táblázatokban és LibreOffice-ban is működnek
(VLOOKUP, INDEX/MATCH, RANK, LARGE, CEILING – nincs dinamikus tömbképlet).
"""

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

from .config import CRITERIA, INDUSTRIES, INPUT_COLUMNS, MISSING_POINTS, PRICING, TIERS

MAX_ROWS = 200  # ennyi ügyfélsorra készülnek elő a képletek
FIRST, LAST = 2, MAX_ROWS + 1

S_IN, S_SET, S_RES, S_RANK, S_HELP = "Ügyfelek", "Beállítások", "Eredmény", "Rangsor", "Útmutató"

FT = '#,##0" Ft"'
PCT = "0%"

HEAD_FILL = PatternFill("solid", fgColor="032D42")
HEAD_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF8E1")
TITLE_FONT = Font(bold=True, size=14, color="032D42")
SECTION_FONT = Font(bold=True, size=12, color="032D42")
THIN = Border(bottom=Side(style="thin", color="C3C2B7"))
TIER_FILLS = {
    "A": PatternFill("solid", fgColor="C8EFC0"),
    "B": PatternFill("solid", fgColor="E3F4D9"),
    "C": PatternFill("solid", fgColor="FDEBC2"),
    "D": PatternFill("solid", fgColor="F6CFCF"),
}


def ref(sheet, cell):
    return f"'{sheet}'!{cell}"


def abs_ref(sheet, column, row):
    return ref(sheet, f"${column}${row}")


def _header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[col(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 45


# ---------------------------------------------------------------- Beállítások

class Layout:
    """A Beállítások lap cellacímei, amikre a képletek hivatkoznak."""

    crit_first = 6
    band_col0 = 8  # H oszloptól jobbra, kritériumonként 3 oszlop
    band_row0 = 7
    ind_first, ind_last = 22, 41
    param_row0 = 45  # hiányzó pont, majd a díjparaméterek
    tier_first = 57

    def __init__(self):
        self.crit_last = self.crit_first + len(CRITERIA) - 1
        self.weight = {c["kod"]: abs_ref(S_SET, "E", self.crit_first + i) for i, c in enumerate(CRITERIA)}
        self.names = f"{ref(S_SET, '$B$%d' % self.crit_first)}:$B${self.crit_last}"
        self.bands = {}
        band_criteria = [c for c in CRITERIA if c["tipus"] == "sav"]
        for i, c in enumerate(band_criteria):
            lo, pt = col(self.band_col0 + 3 * i), col(self.band_col0 + 3 * i + 1)
            last = self.band_row0 + len(c["savok"]) - 1
            self.bands[c["kod"]] = (
                f"{ref(S_SET, f'${lo}${self.band_row0}')}:${pt}${last}",
                abs_ref(S_SET, lo, self.band_row0),
            )
        self.industries = f"{ref(S_SET, f'$A${self.ind_first}')}:$B${self.ind_last}"
        self.industry_names = f"{ref(S_SET, f'$A${self.ind_first}')}:$A${self.ind_last}"
        self.missing = abs_ref(S_SET, "B", self.param_row0)
        self.pricing = {k: abs_ref(S_SET, "B", self.param_row0 + 1 + i) for i, k in enumerate(PRICING)}
        self.tier_last = self.tier_first + len(TIERS) - 1
        self.tiers = f"{ref(S_SET, f'$A${self.tier_first}')}:$E${self.tier_last}"


def build_settings(ws, lay):
    ws["A1"] = "Beállítások – kritériumok, súlyok, sávok, szintek, díjazás"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "A sárga cellák szabadon módosíthatók; minden eredmény azonnal újraszámolódik."

    ws["A4"] = "Kritériumok és súlyok"
    ws["A4"].font = SECTION_FONT
    _header(ws, 5, ["Kód", "Kritérium", "Csoport", "Típus", "Súly", "Leírás"])
    types = {"sav": "sáv", "skala": "1–5 skála", "lista": "lista"}
    for i, c in enumerate(CRITERIA):
        r = lay.crit_first + i
        for j, v in enumerate([c["kod"], c["nev"], c["csoport"], types[c["tipus"]], c["suly"], c["leiras"]], 1):
            ws.cell(row=r, column=j, value=v)
        ws.cell(row=r, column=5).fill = INPUT_FILL
    total = lay.crit_last + 1
    ws.cell(row=total, column=4, value="Összesen").font = Font(bold=True)
    ws.cell(row=total, column=5, value=f"=SUM(E{lay.crit_first}:E{lay.crit_last})").font = Font(bold=True)
    ws.cell(row=total, column=6, value=f'=IF(E{total}=100,"Rendben","Figyelem: a súlyok összege legyen 100!")')
    for w, c in zip([14, 28, 13, 11, 8, 70], "ABCDEF"):
        ws.column_dimensions[c].width = w

    ws.cell(row=4, column=lay.band_col0, value="Sávok (alsó határ → pont)").font = SECTION_FONT
    for i, c in enumerate([c for c in CRITERIA if c["tipus"] == "sav"]):
        c0 = lay.band_col0 + 3 * i
        ws.cell(row=5, column=c0, value=c["nev"]).font = Font(bold=True)
        for j, label in enumerate(["Alsó határ", "Pont"]):
            h = ws.cell(row=6, column=c0 + j, value=label)
            h.fill, h.font = HEAD_FILL, HEAD_FONT
        for k, (lower, pts) in enumerate(c["savok"]):
            for j, v in enumerate([lower, pts]):
                cell = ws.cell(row=lay.band_row0 + k, column=c0 + j, value=v)
                cell.fill = INPUT_FILL
                if j == 0:
                    cell.number_format = "#,##0.##"
        ws.column_dimensions[col(c0)].width = 13
        ws.column_dimensions[col(c0 + 1)].width = 7
        ws.column_dimensions[col(c0 + 2)].width = 3

    ws.cell(row=lay.ind_first - 2, column=1, value="Iparágak pontértéke").font = SECTION_FONT
    for j, label in enumerate(["Iparág", "Pont (1–5)"], 1):
        h = ws.cell(row=lay.ind_first - 1, column=j, value=label)
        h.fill, h.font = HEAD_FILL, HEAD_FONT
    for r in range(lay.ind_first, lay.ind_last + 1):
        for j in (1, 2):
            ws.cell(row=r, column=j).fill = INPUT_FILL
    for i, (name, pts) in enumerate(INDUSTRIES):
        ws.cell(row=lay.ind_first + i, column=1, value=name)
        ws.cell(row=lay.ind_first + i, column=2, value=pts)
    ws.cell(row=lay.ind_first - 2, column=3, value="(új iparág az üres sorokba vehető fel)")

    ws.cell(row=lay.param_row0 - 1, column=1, value="Általános és díjparaméterek").font = SECTION_FONT
    params = [("Hiányzó adat pontja (1–5)", MISSING_POINTS, "0")]
    params += [(label, value, PCT if k == "keret_arany" else FT) for k, (label, value) in PRICING.items()]
    for i, (label, value, fmt) in enumerate(params):
        r = lay.param_row0 + i
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=value)
        cell.fill, cell.number_format = INPUT_FILL, fmt

    ws.cell(row=lay.tier_first - 2, column=1, value="Szintek és díjszorzók").font = SECTION_FONT
    for j, label in enumerate(["Minimum pontszám", "Szint", "Megnevezés", "Díjszorzó", "Jutalék % (később)"], 1):
        h = ws.cell(row=lay.tier_first - 1, column=j, value=label)
        h.fill, h.font = HEAD_FILL, HEAD_FONT
    for i, t in enumerate(TIERS):
        r = lay.tier_first + i
        for j, v in enumerate([t["min"], t["szint"], t["nev"], t["szorzo"], t["jutalek"]], 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.fill = INPUT_FILL
        ws.cell(row=r, column=4).number_format = "0.00"
        ws.cell(row=r, column=5).number_format = "0.0%"
    ws.cell(row=lay.tier_last + 1, column=1,
            value="A szintek a minimum pontszám szerint növekvő sorrendben maradjanak.").font = Font(italic=True)
    ws.freeze_panes = "A6"


# -------------------------------------------------------------------- Ügyfelek

def build_input(ws, lay, clients):
    _header(ws, 1, [h for _, h, _, _ in INPUT_COLUMNS], [w for _, _, w, _ in INPUT_COLUMNS])
    ws.freeze_panes = "B2"

    for i, client in enumerate(clients):
        for j, (key, *_rest) in enumerate(INPUT_COLUMNS, 1):
            value = client.get(key, "")
            ws.cell(row=FIRST + i, column=j, value=None if value == "" else value)

    validations = {
        "lista": DataValidation(type="list", formula1=lay.industry_names, allow_blank=True),
        "igennem": DataValidation(type="list", formula1='"igen,nem"', allow_blank=True),
        "skala": DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True),
        "szam": DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True),
        "penz": DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True),
    }
    validations["skala"].error = "1 és 5 közötti egész számot adj meg."
    for dv in validations.values():
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
    for j, (_, _, _, kind) in enumerate(INPUT_COLUMNS, 1):
        rng = f"{col(j)}{FIRST}:{col(j)}{LAST}"
        if kind in validations:
            validations[kind].add(rng)
        if kind == "penz":
            for r in range(FIRST, LAST + 1):
                ws.cell(row=r, column=j).number_format = FT


def input_cell(key, row):
    idx = [k for k, *_ in INPUT_COLUMNS].index(key) + 1
    return ref(S_IN, f"{col(idx)}{row}")


def criterion_formula(c, lay, row):
    x = input_cell(c["kod"], row)
    miss = lay.missing
    if c["tipus"] == "sav":
        rng, first = lay.bands[c["kod"]]
        body = f"VLOOKUP(MAX({x},{first}),{rng},2,TRUE)"
    elif c["tipus"] == "skala":
        body = f"MIN(5,MAX(1,ROUND({x},0)))"
    else:
        body = f"IFERROR(VLOOKUP(TRIM({x}),{lay.industries},2,FALSE),{miss})"
    return f'=IF($A{row}="","",IF({x}="",{miss},{body}))'


# -------------------------------------------------------------------- Eredmény

RES_FIXED = [
    ("ICP pontszám (0–100)", 12, "0.0"),
    ("Rang", 7, "0"),
    ("Szint", 7, None),
    ("Szint megnevezése", 18, None),
    ("Legerősebb terület", 24, None),
    ("Leggyengébb terület", 24, None),
    ("Munkadíj (Ft/hó)", 14, FT),
    ("Szintszorzó", 10, "0.00"),
    ("Keretarányos minimum (Ft/hó)", 15, FT),
    ("Javasolt havi díj (Ft)", 15, FT),
    ("Díj alapja", 24, None),
    ("Jelenlegi díj (Ft)", 14, FT),
    ("Eltérés (Ft)", 13, FT),
    ("Eltérés (%)", 10, PCT),
    ("Opcionális jutalék (Ft/hó)", 14, FT),
    ("Hiányzó adatok (db)", 10, "0"),
]


class ResCols:
    def __init__(self):
        n = len(CRITERIA)
        self.points = {c["kod"]: col(2 + i) for i, c in enumerate(CRITERIA)}
        names = ["score", "rank", "tier", "tier_name", "best", "worst", "labour", "mult", "budget_min",
                 "fee", "basis", "current", "diff", "diff_pct", "commission", "missing"]
        self.fixed = {k: col(2 + n + i) for i, k in enumerate(names)}
        h0 = 2 + n + len(names) + 1  # egy üres elválasztó oszlop után a segédoszlopok
        self.strength = [col(h0 + i) for i in range(n)]
        self.lost = [col(h0 + n + i) for i in range(n)]
        self.sort_key = col(h0 + 2 * n)
        self.helper_first, self.helper_last = h0, h0 + 2 * n


def build_results(ws, lay, rc):
    headers = ["Ügyfél"] + [c["nev"] + " (pont)" for c in CRITERIA] + [h for h, _, _ in RES_FIXED]
    widths = [26] + [11] * len(CRITERIA) + [w for _, w, _ in RES_FIXED]
    _header(ws, 1, headers, widths)
    helper_headers = [f"erő: {c['kod']}" for c in CRITERIA] + [f"veszt: {c['kod']}" for c in CRITERIA] + ["rendezőkulcs"]
    for i, h in enumerate(helper_headers):
        ws.cell(row=1, column=rc.helper_first + i, value=h).font = Font(italic=True, color="898781")

    f = rc.fixed
    p = lay.pricing
    tiers = lay.tiers
    for r in range(FIRST, LAST + 1):
        blank = f'$A{r}=""'
        ws[f"A{r}"] = f'=IF({input_cell("nev", r)}="","",{input_cell("nev", r)})'
        for c in CRITERIA:
            ws[f"{rc.points[c['kod']]}{r}"] = criterion_formula(c, lay, r)

        score_terms = "+".join(f"{lay.weight[c['kod']]}*({rc.points[c['kod']]}{r}-1)/4" for c in CRITERIA)
        ws[f"{f['score']}{r}"] = f'=IF({blank},"",ROUND({score_terms},1))'
        ws[f"{f['rank']}{r}"] = f'=IF({blank},"",RANK({f["score"]}{r},${f["score"]}${FIRST}:${f["score"]}${LAST}))'
        for key, idx in (("tier", 2), ("tier_name", 3), ("mult", 4)):
            ws[f"{f[key]}{r}"] = f'=IF({blank},"",VLOOKUP({f["score"]}{r},{tiers},{idx},TRUE))'

        s_rng = f"{rc.strength[0]}{r}:{rc.strength[-1]}{r}"
        l_rng = f"{rc.lost[0]}{r}:{rc.lost[-1]}{r}"
        ws[f"{f['best']}{r}"] = f'=IF({blank},"",INDEX({lay.names},MATCH(MAX({s_rng}),{s_rng},0)))'
        ws[f"{f['worst']}{r}"] = (f'=IF({blank},"",IF(MAX({l_rng})=0,"—",'
                                  f'INDEX({lay.names},MATCH(MAX({l_rng}),{l_rng},0))))')
        for i, c in enumerate(CRITERIA):
            pt = f"{rc.points[c['kod']]}{r}"
            ws[f"{rc.strength[i]}{r}"] = f'=IF({blank},"",{pt}+{lay.weight[c["kod"]]}/1000)'
            ws[f"{rc.lost[i]}{r}"] = f'=IF({blank},"",{lay.weight[c["kod"]]}*(5-{pt}))'
        ws[f"{rc.sort_key}{r}"] = f'=IF({blank},"",{f["score"]}{r}-ROW()/100000)'

        yes = lambda key: f'(TRIM({input_cell(key, r)})="igen")'  # noqa: E731
        ws[f"{f['labour']}{r}"] = (
            f'=IF({blank},"",{p["google_alap"]}*{yes("google")}+{p["meta_alap"]}*{yes("meta")}'
            f'+{p["egyeb_alap"]}*N({input_cell("egyeb_csatorna", r)})+{p["oradij"]}*N({input_cell("orak", r)}))'
        )
        ws[f"{f['budget_min']}{r}"] = f'=IF({blank},"",N({input_cell("keret", r)})*{p["keret_arany"]})'
        by_labour = f"{f['labour']}{r}*{f['mult']}{r}"
        raw = f"MAX({p['min_dij']},{by_labour},{f['budget_min']}{r})"
        ws[f"{f['fee']}{r}"] = f'=IF({blank},"",CEILING(ROUND({raw},2),{p["kerekites"]}))'
        ws[f"{f['basis']}{r}"] = (
            f'=IF({blank},"",IF({raw}={by_labour},"Munkaigény × szintszorzó",'
            f'IF({raw}={f["budget_min"]}{r},"Hirdetési keret arányában","Minimumdíj")))'
        )
        cur = input_cell("jelenlegi_dij", r)
        ws[f"{f['current']}{r}"] = f'=IF(OR({blank},{cur}=""),"",{cur})'
        ws[f"{f['diff']}{r}"] = f'=IF({f["current"]}{r}="","",{f["fee"]}{r}-{f["current"]}{r})'
        ws[f"{f['diff_pct']}{r}"] = (f'=IF(OR({f["current"]}{r}="",{f["current"]}{r}=0),"",'
                                     f'{f["diff"]}{r}/{f["current"]}{r})')
        rev = input_cell("bevetel", r)
        ws[f"{f['commission']}{r}"] = f'=IF(OR({blank},{rev}=""),"",{rev}*VLOOKUP({f["score"]}{r},{tiers},5,TRUE))'
        missing = "+".join(f'({input_cell(c["kod"], r)}="")' for c in CRITERIA)
        ws[f"{f['missing']}{r}"] = f'=IF({blank},"",{missing})'

        for (_, _, fmt), key in zip(RES_FIXED, f):
            if fmt:
                ws[f"{f[key]}{r}"].number_format = fmt

    for i in range(rc.helper_first, rc.helper_last + 1):
        ws.column_dimensions[col(i)].hidden = True
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{f['missing']}{LAST}"

    score_rng = f"{f['score']}{FIRST}:{f['score']}{LAST}"
    ws.conditional_formatting.add(score_rng, DataBarRule(start_type="num", start_value=0, end_type="num",
                                                         end_value=100, color="2E981B"))
    _tier_colors(ws, f"{f['tier']}{FIRST}:{f['tier']}{LAST}")
    ws.conditional_formatting.add(
        f"{f['diff']}{FIRST}:{f['diff_pct']}{LAST}",
        CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="1C7A3E", bold=True)))
    ws.conditional_formatting.add(
        f"{f['diff']}{FIRST}:{f['diff_pct']}{LAST}",
        CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C23434", bold=True)))


def _tier_colors(ws, rng):
    top_left = rng.split(":")[0]
    for tier, fill in TIER_FILLS.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{top_left}="{tier}"'], fill=fill))


# --------------------------------------------------------------------- Rangsor

def build_ranking(ws, rc):
    cols = [("Rang", "rank", 7, "0"), ("Ügyfél", None, 28, None), ("ICP pontszám", "score", 12, "0.0"),
            ("Szint", "tier", 7, None), ("Szint megnevezése", "tier_name", 20, None),
            ("Javasolt havi díj", "fee", 16, FT), ("Jelenlegi díj", "current", 15, FT),
            ("Eltérés (Ft)", "diff", 14, FT), ("Eltérés (%)", "diff_pct", 10, PCT),
            ("Legerősebb terület", "best", 24, None), ("Leggyengébb terület", "worst", 24, None),
            ("Hiányzó adatok", "missing", 10, "0")]
    ws["A1"] = "ICP-rangsor – pontszám szerint csökkenő sorrendben"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Automatikusan frissül az Ügyfelek lap alapján. Holtversenynél a korábban felvett ügyfél áll elöl."
    _header(ws, 4, [c[0] for c in cols], [c[2] for c in cols])
    helper = col(len(cols) + 2)
    ws.column_dimensions[helper].hidden = True
    ws[f"{helper}4"] = "sor"
    key_rng = f"{ref(S_RES, f'${rc.sort_key}${FIRST}')}:${rc.sort_key}${LAST}"
    for k in range(1, MAX_ROWS + 1):
        r = 4 + k
        ws[f"{helper}{r}"] = f"=IFERROR(MATCH(LARGE({key_rng},{k}),{key_rng},0),\"\")"
        for j, (_, key, _, fmt) in enumerate(cols, 1):
            src = "A" if key is None else rc.fixed[key]
            src_rng = f"{ref(S_RES, f'${src}${FIRST}')}:${src}${LAST}"
            cell = ws.cell(row=r, column=j, value=f'=IF(${helper}{r}="","",INDEX({src_rng},${helper}{r}))')
            if fmt:
                cell.number_format = fmt
            cell.border = THIN
    ws.freeze_panes = "C5"
    _tier_colors(ws, f"D5:D{4 + MAX_ROWS}")
    ws.conditional_formatting.add(f"C5:C{4 + MAX_ROWS}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=100, color="2E981B"))


# -------------------------------------------------------------------- Útmutató

GUIDE = [
    ("ICP ügyfélminősítő – útmutató", TITLE_FONT),
    ("", None),
    ("Mire jó?", SECTION_FONT),
    ("Az ügyfeleket 11 kritérium alapján 0–100 pontra értékeli, A–D szintbe sorolja, rangsorolja, "
     "és havi díjat javasol az online marketing munkára.", None),
    ("", None),
    ("Használat", SECTION_FONT),
    ("1. Ügyfelek lap: ügyfelenként egy sor. A legördülő listák és az 1–5 skálák ellenőrzöttek. "
     "Az üres mező hiányzó adatnak számít (alapból 2 pont), ezért érdemes mindent kitölteni.", None),
    ("2. Rangsor lap: a pontszám szerinti sorrend, szinttel és díjjavaslattal. Csak olvasható.", None),
    ("3. Eredmény lap: kritériumonkénti pontok és a díjszámítás részletei. Szűrhető.", None),
    ("4. Beállítások lap: a súlyok, sávhatárok, iparági pontok, szintek és díjparaméterek (sárga cellák).", None),
    ("", None),
    ("Pontozás", SECTION_FONT),
    ("Minden kritérium 1–5 pontot kap (sáv, kézi skála vagy iparági lista alapján). "
     "ICP pontszám = Σ súly × (pont − 1) / 4, így 0 és 100 közé esik.", None),
    ("Szintek: A ≥ 80 (ideális ügyfél), B ≥ 65, C ≥ 50, D < 50.", None),
    ("A „Leggyengébb terület” az a kritérium, ahol a legtöbb súlyozott pont veszik el – "
     "itt érdemes fejleszteni vagy tárgyalni.", None),
    ("", None),
    ("Díjjavaslat", SECTION_FONT),
    ("Munkadíj = csatornák alapdíja (Google, Meta, egyéb) + becsült havi óra × óradíj.", None),
    ("Javasolt díj = a három közül a legnagyobb, felfelé kerekítve: "
     "munkadíj × szintszorzó | hirdetési keret × keretarány | minimumdíj.", None),
    ("A szintszorzó a kockázatot és a többletenergiát árazza: A = 1,00; B = 1,05; C = 1,15; D = 1,30.", None),
    ("Az „Eltérés” a jelenlegi díjhoz képest mutatja, hol érdemes árat emelni (pozitív) "
     "vagy hol fizet az ügyfél a javasoltnál többet (negatív).", None),
    ("Jutalék: a Beállítások „Jutalék % (később)” oszlopa most 0. Ha kitöltöd, a mért bevétel "
     "arányában számol opcionális jutalékot.", None),
    ("", None),
    ("Google Táblázatok", SECTION_FONT),
    ("A fájl feltölthető Google Drive-ra és megnyitható Google Táblázatként; a képletek ott is működnek.", None),
]


def build_guide(ws):
    ws.column_dimensions["A"].width = 120
    for i, (text, font) in enumerate(GUIDE, 1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            cell.font = font


def build_workbook(clients=()):
    wb = Workbook()
    guide = wb.active
    guide.title = S_HELP
    ws_in = wb.create_sheet(S_IN)
    ws_rank = wb.create_sheet(S_RANK)
    ws_res = wb.create_sheet(S_RES)
    ws_set = wb.create_sheet(S_SET)

    lay, rc = Layout(), ResCols()
    build_guide(guide)
    build_settings(ws_set, lay)
    build_input(ws_in, lay, list(clients))
    build_results(ws_res, lay, rc)
    build_ranking(ws_rank, rc)

    for ws, color in ((ws_in, "FAB219"), (ws_rank, "2E981B"), (ws_res, "134F6C"), (ws_set, "898781")):
        ws.sheet_properties.tabColor = color
    wb.active = 1
    return wb
