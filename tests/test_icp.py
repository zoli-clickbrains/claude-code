import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from icp.config import CRITERIA, default_config
from icp.engine import criterion_points, score_client, score_clients
from icp.sample import SAMPLE_CLIENTS
from icp.workbook import FIRST, S_RANK, S_RES, ResCols, build_workbook

CFG = default_config()
BY_CODE = {c["kod"]: c for c in CRITERIA}


def test_weights_sum_to_100():
    assert sum(c["suly"] for c in CRITERIA) == 100


@pytest.mark.parametrize("value,expected", [(0, 1), (299_999, 1), (300_000, 2), (2_999_999, 3), (10_000_000, 5)])
def test_budget_bands(value, expected):
    assert criterion_points(BY_CODE["keret"], value, CFG) == expected


@pytest.mark.parametrize("days,expected", [(0, 5), (1, 4), (7, 4), (8, 3), (30, 2), (31, 1), (90, 1)])
def test_payment_delay_bands_are_descending(days, expected):
    assert criterion_points(BY_CODE["keses"], days, CFG) == expected


def test_missing_and_unknown_values_get_missing_points():
    assert criterion_points(BY_CODE["roas"], "", CFG) == CFG["missing_points"]
    assert criterion_points(BY_CODE["iparag"], "Ismeretlen", CFG) == CFG["missing_points"]
    assert criterion_points(BY_CODE["iparag"], " e-kereskedelem, webshop ", CFG) == 5


def test_perfect_and_worst_clients_hit_score_limits():
    best = {c["kod"]: 5 for c in CRITERIA if c["tipus"] == "skala"}
    best.update(nev="X", iparag="E-kereskedelem, webshop", arbevetel=5000, keret=20_000_000, roas=10, keses=0,
                hossz=36)
    r = score_client(best)
    assert (r.pontszam, r.szint, r.leggyengebb) == (100.0, "A", "—")

    worst = {c["kod"]: 1 for c in CRITERIA if c["tipus"] == "skala"}
    worst.update(nev="Y", iparag="Gyártás, ipar", arbevetel=0, keret=0, roas=0, keses=60, hossz=0)
    r = score_client(worst)
    assert r.szint == "D" and r.pontszam < 10


def test_fee_uses_largest_of_three_components():
    base = {"nev": "Z", "google": "igen", "meta": "nem", "orak": 10, "keret": 0}
    r = score_client(base)
    assert r.munkadij == 60_000 + 10 * 12_000
    assert r.dij_alapja == "Munkaigény × szintszorzó"
    assert r.javasolt_dij == 240_000  # 180 000 × 1,30 = 234 000 → felfelé kerekítve

    r = score_client({**base, "keret": 8_000_000})
    assert (r.javasolt_dij, r.dij_alapja) == (640_000, "Hirdetési keret arányában")

    r = score_client({"nev": "Kicsi", "google": "nem", "meta": "nem", "orak": 1})
    assert (r.javasolt_dij, r.dij_alapja) == (150_000, "Minimumdíj")


def test_ranking_orders_and_handles_ties():
    results = score_clients([{"nev": "B"}, {"nev": "A"}, {"nev": ""}])
    assert [r.nev for r in results] == ["B", "A"]
    assert [r.rang for r in results] == [1, 1]


# ---- A táblázat képletei ugyanazt adják, mint a Python-motor ----

@pytest.fixture(scope="module")
def recalculated(tmp_path_factory):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pytest.skip("LibreOffice nincs telepítve")
    d = tmp_path_factory.mktemp("wb")
    src = d / "icp.xlsx"
    build_workbook(SAMPLE_CLIENTS).save(src)
    out = d / "out"
    subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{d}/lo", "--convert-to", "xlsx",
                    "--outdir", str(out), str(src)], check=True, capture_output=True, timeout=180)
    return load_workbook(out / "icp.xlsx", data_only=True)


def test_workbook_formulas_match_engine(recalculated):
    ws = recalculated[S_RES]
    rc = ResCols()
    f = rc.fixed
    expected = {r.nev: r for r in score_clients(SAMPLE_CLIENTS)}
    for i in range(len(SAMPLE_CLIENTS)):
        row = FIRST + i
        name = ws[f"A{row}"].value
        e = expected[name]
        for c in CRITERIA:
            assert ws[f"{rc.points[c['kod']]}{row}"].value == e.pontok[c["kod"]], (name, c["kod"])
        got = {k: ws[f"{f[k]}{row}"].value for k in f}
        assert got["score"] == pytest.approx(e.pontszam), name
        assert got["rank"] == e.rang
        assert (got["tier"], got["tier_name"]) == (e.szint, e.szint_nev)
        assert (got["best"], got["worst"]) == (e.legerosebb, e.leggyengebb), name
        assert got["fee"] == e.javasolt_dij, name
        assert got["basis"] == e.dij_alapja, name
        assert got["missing"] == e.hianyzo_adatok
        assert (got["diff"] if got["diff"] != "" else None) == e.elteres
    assert ws[f"A{FIRST + len(SAMPLE_CLIENTS)}"].value in (None, "")


def test_ranking_sheet_is_sorted(recalculated):
    ws = recalculated[S_RANK]
    names = [ws[f"B{r}"].value for r in range(5, 5 + len(SAMPLE_CLIENTS))]
    assert names == [r.nev for r in score_clients(SAMPLE_CLIENTS)]
    assert ws[f"B{5 + len(SAMPLE_CLIENTS)}"].value in (None, "")
